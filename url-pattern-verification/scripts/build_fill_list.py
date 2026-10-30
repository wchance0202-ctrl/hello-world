#!/usr/bin/env python3
"""Build the fill-in list for the user's summary table (站点名称｜pattern链接｜黑白名单属性｜正则表达式).

Sheet 1 "待补正则": every path entry of both batches that has no regex at all (no formal,
candidate or failed rule), with why it is empty and the clue the last round saw: the most
common URL structures of the same-site links on the entry's own pages that fell outside the
entry path (that is where the column's items usually live), with counts and sample URLs.
Sheet 2 "已填需复核": rules that may already be in the user's table but must be fixed or
re-checked (scope too wide, failed, candidate).

Usage: python3 -I build_fill_list.py <out_dir>
"""
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
BATCHES = (("第一批", ROOT / "output/round3"), ("第二批", ROOT / "batch2/output/round3"))
FILL_COLS = ["pattern链接", "entry_id", "原表行", "空白原因", "建议做法", "上一轮线索：入口页上前缀外同站链接的主要结构（数量｜样例）",
             "当前发现覆盖", "当前链接状态"]
REVIEW_COLS = ["pattern链接", "entry_id", "原表行", "表中可能已有的正则", "问题", "怎么改"]
WHY = {
    "blocked": ("入口打不开（拦截 / 超时 / 证书 / 代理 / 挑战页）", "以正常身份打开后按全流程处理；仍打不开就留空并写访问受限"),
    "outside": ("条目不在入口路径下（链接指向同站其他路径）", "看线索列里的结构，按第 2 节做“专属性证明”；能证明就写正则，不能就写精确清单或留空"),
    "sibling": ("同结构的链接是同级栏目页，不是本栏目的条目", "排除同级栏目页，找本栏目真正的条目地址（常在前缀外），再做专属性证明"),
    "small": ("同结构的条目只有 1–3 个，不够归纳", "核对是否真的只有这几个；是的话写精确清单正则（快照）"),
    "nolinks": ("页面没有标准链接，内容由脚本渲染", "浏览器点击卡片或看页面自己的只读数据请求，取得条目地址后再归纳"),
    "hash": ("Hash 路由，只取到首屏", "浏览器翻完列表、点击取样，逐个单独打开核验"),
    "none": ("入口页上没有任何本栏目条目链接", "确认是否是单页、活动页或下载页；是的话按第 5 节留空或写精确清单"),
    "other": ("上一轮有家族但没有形成规则", "按全流程重做，注意尾斜杠、共享路径和页面隐藏的条目"),
}


def struct(u):
    p = urlsplit(u)
    segs = [s for s in p.path.split("/") if s]
    head = "/".join(segs[:-1])
    key = f"{p.netloc}/{head + '/' if head else ''}{{x}}{'/' if p.path.endswith('/') and segs else ''}"
    if p.query:
        key += "?" + "&".join(kv.split("=", 1)[0] + "={…}" for kv in p.query.split("&"))
    return key


def clue(urls, limit=4):
    by = defaultdict(list)
    for u in urls:
        by[struct(u)].append(u)
    top = sorted(by.items(), key=lambda kv: -len(kv[1]))[:limit]
    return "\n".join(f"{k} ×{len(v)}｜{v[0]}" + (f"；{v[-1]}" if len(v) > 1 else "") for k, v in top)


def main(out):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    fill, review = [], []
    for batch, d in BATCHES:
        E = json.loads((d / "entries.json").read_text(encoding="utf-8"))
        R = json.loads((d / "rules.json").read_text(encoding="utf-8"))
        F = {f["entry_id"]: f for f in json.loads((d / "families.json").read_text(encoding="utf-8"))}
        has, kept = set(), defaultdict(list)
        for k in ("formal", "candidates", "failed"):
            for r in R[k]:
                if r.get("pattern"):
                    has.update(r.get("entry_ids", []))
        for r in R["kept_as_links"]:
            for eid in r.get("entry_ids", []):
                kept[eid].append(r)
        byid = {e["entry_id"]: e for e in E}
        for e in E:
            if e["kind"] != "path_entry" or e["entry_id"] in has:
                continue
            eid = e["entry_id"]
            f = F.get(eid, {})
            text = (e.get("round3") or {}).get("discovery_text", "").split("\n")[0]
            if e.get("discovery_status") == "未验证":
                k = "blocked"
            elif text.startswith("部分：Hash"):
                k = "hash"
            elif "没有可枚举的标准链接" in text:
                k = "nolinks"
            elif eid in kept:
                k = "sibling" if any("同级" in json.dumps(r, ensure_ascii=False) for r in kept[eid]) else "small"
            elif f.get("outside_prefix"):
                k = "outside"
            elif "未发现可归纳" in text:
                k = "none"
            else:
                k = "other"
            fill.append({"pattern链接": e["raw"], "entry_id": eid, "原表行": f"{batch} {e['row']}",
                         "空白原因": WHY[k][0], "建议做法": WHY[k][1],
                         "上一轮线索：入口页上前缀外同站链接的主要结构（数量｜样例）": clue(f.get("outside_prefix", [])),
                         "当前发现覆盖": e.get("discovery_status", ""), "当前链接状态": e.get("link_status", "")})
        for kind in ("formal", "failed", "candidates"):
            for r in R[kind]:
                if not r.get("pattern"):
                    continue
                if kind == "formal" and " 入口前缀 / 下" not in r.get("target", ""):
                    continue
                eid = r["entry_ids"][0]
                if kind == "formal":
                    prob, fix = ("范围超界：入口在站点根目录，正则按整个站点泛化，会匹配其他栏目的页面",
                                 "找出栏目标识（如 catid=20、type=1）并写成字面量；条目地址不带标识时改为精确清单")
                elif kind == "failed":
                    prob, fix = ("验证不通过：" + "；".join(r.get("why", []) or r.get("why_not_verified", []))[:200]
                                 + "（规则会匹配同级栏目页）", "删掉这条，按待补正则的流程重做")
                else:
                    prob, fix = ("候选，未达准入：" + "；".join(r.get("why", []) or r.get("why_not_verified", []))[:200],
                                 "补齐所缺证据后保留，或修正")
                review.append({"pattern链接": byid[eid]["raw"], "entry_id": eid, "原表行": f"{batch} {byid[eid]['row']}",
                               "表中可能已有的正则": r["pattern"], "问题": prob, "怎么改": fix})
    for name, cols, rows in (("fill-list", FILL_COLS, fill), ("review-list", REVIEW_COLS, review)):
        with open(out / f"{name}.csv", "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)
    wb = Workbook()
    for i, (title, cols, rows, widths) in enumerate((
            ("待补正则", FILL_COLS, fill, (55, 13, 12, 34, 50, 90, 24, 30)),
            ("已填需复核", REVIEW_COLS, review, (55, 13, 12, 75, 60, 50)))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for r in rows:
            ws.append([r[c] for c in cols])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFF2CC")
        for col, wd in zip("ABCDEFGH", widths):
            ws.column_dimensions[col].width = wd
        for rr in ws.iter_rows(min_row=2):
            for c in rr:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "B2"
    wb.save(out / "fill-list.xlsx")
    print(json.dumps({"fill": len(fill), "fill_by_reason": Counter(r["空白原因"] for r in fill),
                      "review": len(review), "review_by_problem": Counter(r["问题"][:6] for r in review)},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main(sys.argv[1])
