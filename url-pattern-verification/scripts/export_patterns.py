#!/usr/bin/env python3
"""Export the usable rules of one batch into patterns-final.xlsx|csv.

Sheet 1 "可用正则": verified rules plus the ones the user adopted (marked), one row per rule,
with the entry it belongs to and a landing-checked match example.
Sheet 2: candidates and failed rules (for reference only, not for use).

Usage: python3 -I export_patterns.py <round3_dir>
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLS = ["原表行", "entry_id", "入口URL", "Pattern 正则", "状态", "匹配样例", "规则ID", "说明"]


def main(r3):
    r3 = Path(r3)
    R = json.loads((r3 / "rules.json").read_text(encoding="utf-8"))
    E = {e["entry_id"]: e for e in json.loads((r3 / "entries.json").read_text(encoding="utf-8"))}

    def row(r, usable):
        eid = r["entry_ids"][0]
        e = E[eid]
        s = r.get("samples", {})
        ex = s.get("holdout_pos", [])[:1] or s.get("inference", [])[:1]
        note = ""
        if not usable or r.get("adopted_by_user"):
            note = "；".join(r.get("why", []) or r.get("why_not_verified", []))[:300]
        elif r.get("risks"):
            note = "；".join(r["risks"])
        return {"原表行": e["row"], "entry_id": eid, "入口URL": e["raw"], "Pattern 正则": r.get("pattern", ""),
                "状态": r["status"], "匹配样例": ex[0] if ex else "", "规则ID": r["id"], "说明": note}

    use = sorted((row(r, True) for r in R["formal"]), key=lambda x: (x["原表行"], x["entry_id"], x["规则ID"]))
    other = sorted((row(r, False) for r in R["candidates"] + R["failed"] if "entry_ids" in r and r.get("pattern")),
                   key=lambda x: (x["原表行"], x["entry_id"], x["规则ID"]))
    with open(r3 / "patterns-final.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        w.writerows(use)
    wb = Workbook()
    for i, (title, data) in enumerate((("可用正则", use), ("候选与未通过（不建议直接用）", other))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title[:31]
        ws.append(COLS)
        for r in data:
            ws.append([r[c] for c in COLS])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="FFF2CC")
        for col, wd in zip("ABCDEFGH", (8, 13, 55, 75, 22, 60, 22, 60)):
            ws.column_dimensions[col].width = wd
        for rr in ws.iter_rows(min_row=2):
            for c in rr:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            rr[3].font = Font(name="Consolas")
        ws.freeze_panes = "A2"
    wb.save(r3 / "patterns-final.xlsx")
    print(json.dumps({"usable": len(use), "entries": len({r["entry_id"] for r in use}), "other": len(other)}))


if __name__ == "__main__":
    main(sys.argv[1])
