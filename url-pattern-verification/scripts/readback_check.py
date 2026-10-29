#!/usr/bin/env python3
"""Full read-back of result-table.xlsx against the input and entries.json."""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

FIELDS = ["Pattern 正则", "保留链接", "匹配样例", "结果类型", "规则验证", "来源映射", "发现覆盖", "备注与风险"]


def main(inp, out):
    body = [l for l in Path(inp).read_text(encoding="utf-8").split("\n")[1:] if l != ""]
    ents = json.loads((out / "entries.json").read_text(encoding="utf-8"))
    rules = json.loads((out / "rules.json").read_text(encoding="utf-8"))
    ws = load_workbook(out / "result-table.xlsx").active
    hdr = [c.value for c in ws[1]]
    errs = []
    if hdr != ["指定pattern"] + FIELDS:
        errs.append(f"header {hdr}")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    expected_rows = len({e["row"] for e in ents})
    if len(rows) != expected_rows:
        errs.append(f"rows {len(rows)} != {expected_rows}")
    # concatenating the cells in order must reproduce the input lines exactly
    joined = [l for r in rows for l in (r[0] or "").split("\n")]
    if joined != body:
        errs.append("original values differ from input")
    by_row = {}
    for e in ents:
        by_row.setdefault(e["row"], []).append(e)
    maxlen = 0
    for i, r in enumerate(rows, 2):
        vals = dict(zip(FIELDS, r[1:]))
        maxlen = max([maxlen] + [len(v) for v in r if isinstance(v, str)])
        kinds = {e["kind"] for e in by_row[i]}
        if kinds == {"placeholder"} and any(vals.values()):
            errs.append(f"R{i} placeholder has results")
        if kinds == {"host_exempt"} and (vals["结果类型"] != "host免提" or
                                         any(v for k, v in vals.items() if k != "结果类型")):
            errs.append(f"R{i} host row not host-only")
        for e in by_row[i]:
            if e["kind"] == "path_entry":
                for f in ("保留链接", "来源映射", "发现覆盖"):
                    if f"{e['entry_id']}" not in (vals[f] or "") or (f != "发现覆盖" and e["raw"] not in vals[f]):
                        errs.append(f"{e['entry_id']} missing in {f}")
        if vals["Pattern 正则"]:
            formal = {x["pattern"] for x in rules["formal"]}  # verified + user-adopted rules
            if any(p not in formal for p in vals["Pattern 正则"].split("\n")):
                errs.append(f"R{i} pattern not in formal rules")
    print(json.dumps({"rows": len(rows), "max_cell_chars": maxlen, "excel_cell_limit": 32767,
                      "errors": errs, "ok": not errs and maxlen < 32767}, ensure_ascii=False))
    return 0 if not errs else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], Path(sys.argv[2])))
