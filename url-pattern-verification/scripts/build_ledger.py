#!/usr/bin/env python3
"""Build the entry ledger, result table, resume queue and host allowlist.

Offline step: no network. Reads the inline column text (one item per line,
header first) plus the skill classifier output, reconstructs the table
rows, and writes every deliverable that does not depend on site access.

Usage:
  python3 -I build_ledger.py <input_inline.txt> <classified.json> <out_dir>
"""
import csv
import json
import sys
from collections import OrderedDict
from pathlib import Path

import os

SHEET = os.environ.get("LEDGER_SHEET", "S1")  # sheet id prefix of every entry_id (S1, S2, ...)
COLUMN = "指定pattern"
RUN_DATE = "2026-10-26"
NET_BLOCK = ("本轮(2026-10-26)未访问：会话网络策略对目标站 CONNECT 返回 403"
             "（代理 connect_rejected），WebFetch 亦无法解析目标域名")
RESULT_FIELDS = ["Pattern 正则", "保留链接", "匹配样例", "结果类型", "规则验证",
                 "来源映射", "发现覆盖", "备注与风险"]
ROW_NOTE = ("行边界为重建：内联粘贴丢失引号，按连续同注册域分组得 152 行"
            "（107 个 URL 单元格+45 个“/”行），与前轮检查点计数一致，待与原始文件核对")  # sheet S1


def row_note(n_rows, n_cells, n_skip):
    if SHEET == "S1":
        return ROW_NOTE
    return (f"行边界为重建：内联粘贴丢失引号，按连续同注册域分组得 {n_rows} 行"
            f"（{n_cells} 个 URL 单元格+{n_skip} 个“/”行），待与原始文件核对")

# Previous-round facts the user reported (2026-10-23 checkpoint). The evidence
# ZIP was not supplied, so these are carried as history, never as verified.
PREV_403 = {
    "https://www.greenbook.org/directory",
    "https://www.greenbook.org/insights",
    "https://www.greenbook.org/grit",
    "https://www.cbinsights.com/research/",
    "https://www.cbinsights.com/strategy-terminal/",
    "https://dpaaglobal.com/omnichannel-study/",
}

KNOWN_ISSUES = {
    "https://www.forrester.com/blogs/category/b2b-research/": [
        ("K1", "前轮：首页与第2页返回相同20个标题链接（16个唯一详情）；详情位于共享 /blogs/{slug}/，"
               "不能用 /blogs/[slug]/ 作为 B2B Research 专属规则；前轮仅有已确认成员的有限集合候选，"
               "不覆盖未来文章。需查清真实分页机制、页码是否推进、是否缓存/动态请求。"),
    ],
    "https://www.forrester.com/blogs/category/b2b-research/page/2/": [
        ("K1", "前轮：与第1页内容相同，不能当作新增内容；需核对是否服务器忽略页码/缓存/动态加载。"),
    ],
    "https://www.numerator.com/resources/": [
        ("K2", "前轮：列表经只读 POST https://www.numerator.com/wp-admin/admin-ajax.php "
               "action=filter_resources（searchText、categoryText、postperpage、currentPostCount、"
               "insightsPaged、type），空关键词+全部分类每页8条，取得606个唯一资源链接并与总数606对账，"
               "第77次响应返回 false 结束标志。仅说明当时该筛选条件下列表枚举完成；606个详情/文件/外站跳转未逐项验收。"
               "恢复时先确认总量与筛选条件是否变化，不全量重抓。URL 正则无法表达 POST body 条件。"),
    ],
    "https://contentsquare.com/blog/": [
        ("K3", "前轮：静态HTML不足，需浏览器渲染；已访问32个DOM列表页、浏览器分页队列耗尽；"
               "筛选分支、全部详情与跳转未验收。"),
    ],
    "https://yougov.com/content/economy": [
        ("K4", "前轮：部分详情落在共享 /articles/ 路径，不能仅凭通用路径认定栏目成员；业务脚本域 "
               "d18w79e3uiqk2z.cloudfront.net 曾被拦截，后单脚本200，但完整动态加载与筛选未重新验收。"
               "本轮该域同样被网络策略拒绝。"),
    ],
    "https://www.gartner.com/en/marketing/insights": [
        ("K5", "前轮跳转样本：部分 /insights/ 链接跳到 /research/、/articles/ 等，不被原候选覆盖；按真实跳转链补规则或保留具体目标。"),
    ],
    "https://www.mintel.com/insights/": [
        ("K5", "前轮跳转样本：部分 insights 链接跳到 /trend-predictions/，不被原候选覆盖。"),
    ],
    "https://www.optimizely.com/field-notes/guides": [
        ("K5", "前轮跳转样本：部分页面去掉尾斜杠，最终地址不被原候选覆盖；需按真实跳转链分别处理。"),
    ],
    "https://www.optimizely.com/field-notes/articles": [
        ("K5", "前轮跳转样本：部分页面去掉尾斜杠，最终地址不被原候选覆盖；需按真实跳转链分别处理。"),
    ],
}
for _u in ("https://yougov.com/content/consumer", "https://yougov.com/reports",
           "https://yougov.com/business/sectors/consumer"):
    KNOWN_ISSUES[_u] = KNOWN_ISSUES["https://yougov.com/content/economy"]
for _u in ("https://www.gartner.com/en/sales/insights",
           "https://www.gartner.com/en/customer-service-support/insights",
           "https://www.gartner.com/en/finance/insights",
           "https://www.gartner.com/en/insights/finance"):
    KNOWN_ISSUES[_u] = KNOWN_ISSUES["https://www.gartner.com/en/marketing/insights"]
# WIPO: the user did not name which sub-pages redirect to .../index, so the note
# applies to every www.wipo.int/en/web/ entry until the redirect chain is re-read.
WIPO_K5 = [("K5", "前轮跳转样本：WIPO 部分子页面跳到附带 /index 的路径（前轮未指明具体页面，本轮未能复核）。")]

KNOWN_DEP_HOSTS = {
    "d18w79e3uiqk2z.cloudfront.net": "YouGov 业务脚本（前轮观察）",
}


def reg_domain(host):
    parts = host.split(".")
    if len(parts) >= 3 and parts[-2] in ("co", "com", "org") and len(parts[-1]) == 2:
        return ".".join(parts[-3:])
    return ".".join(parts[-2:])


def reconstruct_rows(items):
    """Consecutive same-registrable-domain URLs form one cell; each '/' is a row."""
    rows, prev = [], None
    for i, it in enumerate(items):
        if it["kind"] == "skip":
            rows.append([i])
            prev = None
            continue
        rd = reg_domain(it["host"])
        if rd == prev:
            rows[-1].append(i)
        else:
            rows.append([i])
            prev = rd
    return rows


def entry_notes(raw):
    notes = list(KNOWN_ISSUES.get(raw, []))
    if raw.startswith("https://www.wipo.int/en/web/"):
        notes += WIPO_K5
    return notes


def build(input_path, classified_path, out):
    lines = [l for l in Path(input_path).read_text(encoding="utf-8").split("\n")]
    header, body = lines[0], [l for l in lines[1:] if l != ""]
    assert header == COLUMN, header
    items = json.loads(Path(classified_path).read_text(encoding="utf-8"))
    assert [it["raw"] for it in items] == body
    rows = reconstruct_rows(items)

    entries, table, queue = [], [], []
    for r_idx, idxs in enumerate(rows):
        row_no = r_idx + 2  # header occupies row 1
        cell_raw = "\n".join(items[i]["raw"] for i in idxs)
        cell = {"原表行": row_no, COLUMN: cell_raw}
        path_entries, host_entries = [], []
        for e_no, i in enumerate(idxs, 1):
            it = items[i]
            eid = f"{SHEET}-R{row_no}-E{e_no}"
            kind = {"entry": "path_entry", "missing_scheme": "path_entry", "host": "host_exempt",
                    "skip": "placeholder"}[it["kind"]]
            rec = {
                "entry_id": eid, "sheet": SHEET, "row": row_no, "entry_index": e_no,
                "line_index": i + 2, "raw": it["raw"], "kind": kind,
                "host": it.get("host"), "path": it.get("path"), "query": it.get("query"),
                "fragment": it.get("fragment"),
                "row_boundary": "reconstructed",
                # a path without scheme keeps its raw value; it is requested as https:// first and the
                # scheme that actually worked is recorded by the probe
                **({"request_url": "https://" + it["raw"], "scheme_assumed": "https（待实测）"}
                   if it["kind"] == "missing_scheme" else {}),
            }
            if kind == "path_entry":
                rec.update({
                    "discovery_status": "未验证",
                    "rule_status": "无规则",
                    "link_status": "待核验",
                    "access_this_round": NET_BLOCK,
                    "previous_checkpoint": (
                        "2026-10-23 前轮：HTTP 403（历史，不作永久结论）" if it["raw"] in PREV_403
                        else "2026-10-23 前轮：汇总称 HTTP 200 / 部分完成；逐入口证据未随附"),
                    "known_issues": [{"id": k, "text": t} for k, t in entry_notes(it["raw"])],
                })
                path_entries.append(rec)
            elif kind == "host_exempt":
                host_entries.append(rec)
            entries.append(rec)

        fields = {f: "" for f in RESULT_FIELDS}
        if path_entries or host_entries:
            types = []
            if path_entries:
                types.append(f"{_span(path_entries)} 链接保留（入口保留形式）")
            if host_entries:
                types.append(f"{_span(host_entries)} host免提")
            fields["结果类型"] = "；".join(types) if len(idxs) > 1 else types[0].split(" ", 1)[1]
            if not path_entries:  # pure host cell: only the result type
                fields["结果类型"] = "host免提"
        if path_entries:
            fields["保留链接"] = "\n".join(f"{e['entry_id']}: {e['raw']}" for e in path_entries)
            fields["规则验证"] = (
                "正式规则 0；候选 0（前轮 255 条候选未随附，不能复用或宣称）。"
                f"\n{NET_BLOCK}；页面身份、结构、范围均未核验。"
                "\n引擎就绪：google-re2 1.1.20251105（Python）、node-re2 1.24.1；本轮无规则测试对。")
            fields["来源映射"] = "\n".join(
                f"{e['entry_id']} = {SHEET}!{COLUMN} 第{e['row']}行 第{e['entry_index']}个入口: {e['raw']}"
                for e in path_entries)
            fields["发现覆盖"] = "\n".join(
                f"{e['entry_id']}: 未验证（本轮0请求）；{e['previous_checkpoint']}；续跑见 queue.json"
                for e in path_entries)
            risks = [row_note(len(rows), sum(1 for r in rows if items[r[0]]["kind"] != "skip"),
                              sum(1 for r in rows if items[r[0]]["kind"] == "skip"))]
            for e in path_entries:
                for k in e["known_issues"]:
                    risks.append(f"{e['entry_id']} [{k['id']}] {k['text']}")
            fields["备注与风险"] = "\n".join(risks)
            for e in path_entries:
                queue.extend(queue_items(e))
        table.append({**cell, **fields})

    hosts = OrderedDict()
    for e in entries:
        if e["kind"] == "path_entry":
            hosts.setdefault(e["host"], []).append(e["entry_id"])

    out.mkdir(parents=True, exist_ok=True)
    _dump(out / "entries.json", entries)
    _dump(out / "queue.json", queue)
    _dump(out / "rules.json", {
        "target_engine": "re2", "engines_ready": [
            {"name": "google-re2", "language": "python", "version": "1.1.20251105"},
            {"name": "node-re2", "language": "node", "version": "1.24.1",
             "harness": "extract-url-patterns/scripts/validate_patterns.mjs"}],
        "generated_at": RUN_DATE, "formal": [], "candidates": [],
        "note": ("本轮无可交付规则：目标站全部被网络策略拒绝；前轮 255 条候选与证据包未随附。"
                 "formal 仅收录满足技能准入的已验证规则。")})
    with open(out / "hosts-allowlist.txt", "w", encoding="utf-8") as f:
        f.write("# 续跑需放行的目标域名（路径入口所在 host，不含 24 个 host 免提项）\n")
        for h, ids in hosts.items():
            f.write(f"{h}\t# {len(ids)} 个入口: {', '.join(ids)}\n")
        f.write("# 已知依赖域（前轮观察）\n")
        for h, why in KNOWN_DEP_HOSTS.items():
            f.write(f"{h}\t# {why}\n")
        f.write("# 动态页面还会加载未知第三方脚本/接口域；浏览器渲染类入口建议使用完整网络访问。\n")
    write_table(out, table)
    return entries, table, queue, hosts


def _span(es):
    if len(es) == 1:
        return f"E{es[0]['entry_index']}"
    nums = [e["entry_index"] for e in es]
    if nums == list(range(nums[0], nums[-1] + 1)):
        return f"E{nums[0]}–E{nums[-1]}"
    return ",".join(f"E{n}" for n in nums)


def queue_items(e):
    eid, raw = e["entry_id"], e["raw"]
    base = [
        ("Q1-access", "请求入口，保存逐跳跳转、最终地址、状态码、标题/正文身份；区分软404/挑战页/空壳"),
        ("Q2-discover", "建立 入口→栏目/分支→分页→主卡片→详情→跳转/文件 账本；区分导航/页脚/推荐；普通锚点不计页"),
        ("Q3-paginate", "推进分页/加载更多直至结束标志或已证实无新增；比较各页唯一详情集合，识别忽略页码/缓存"),
        ("Q4-dynamic", "静态不足时检查渲染DOM、业务脚本、只读接口与点击路由"),
        ("Q5-family", "按 URL 家族决定泛化/有限集合/保留；共享路径不泛化"),
        ("Q6-validate", "冻结独立样本：每条泛化规则≥3归纳正例、≥1独立正例、≥3同站近邻独立反例；RE2 实测"),
    ]
    items = [{"queue_id": f"{eid}-{qid}", "entry_id": eid, "url": raw, "action": text,
              "reason": "本轮网络策略拒绝目标站", "parent": None, "cursor": None,
              "priority": 2, "status": "pending"} for qid, text in base]
    if raw in PREV_403:
        items[0]["action"] += "；前轮 403，按当前实际访问重新判断，不沿用历史结论"
        items[0]["priority"] = 1
    for k in e["known_issues"]:
        items.append({"queue_id": f"{eid}-{k['id']}", "entry_id": eid, "url": raw,
                      "action": k["text"], "reason": "前轮已发现问题，须保留并复核",
                      "parent": None, "cursor": None, "priority": 1, "status": "pending"})
    return items


def write_table(out, table):
    cols = ["原表行", COLUMN] + RESULT_FIELDS
    with open(out / "result-table.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(table)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "结果表"
    ws.append([COLUMN] + RESULT_FIELDS)  # original header in A1, row numbers preserved
    for r in table:
        ws.append([r[COLUMN]] + [r[f] for f in RESULT_FIELDS])
    hdr_font = Font(bold=True)
    for c in ws[1]:
        c.font = hdr_font
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1][1:]:
        c.fill = PatternFill("solid", fgColor="FFF2CC")
    widths = {"A": 70, "B": 40, "C": 60, "D": 20, "E": 34, "F": 60, "G": 80, "H": 70, "I": 100}
    for col, wd in widths.items():
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[1].font = Font(name="Consolas")
    ws.freeze_panes = "B2"
    wb.save(out / "result-table.xlsx")


def _dump(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


if __name__ == "__main__":
    ents, tab, q, hs = build(sys.argv[1], sys.argv[2], Path(sys.argv[3]))
    kinds = {}
    for e in ents:
        kinds[e["kind"]] = kinds.get(e["kind"], 0) + 1
    print(json.dumps({"rows": len(tab), "items": len(ents), "kinds": kinds,
                      "unique_path_urls": len({e["raw"] for e in ents if e["kind"] == "path_entry"}),
                      "hosts_to_allow": len(hs), "queue_items": len(q)}, ensure_ascii=False))
